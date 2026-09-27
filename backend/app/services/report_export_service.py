"""Report export: Word (docxtpl) and PDF (WeasyPrint + Jinja2)."""

from __future__ import annotations

import re
from abc import ABC, abstractmethod
from datetime import UTC, datetime
from html import unescape
from pathlib import Path
from typing import Any

from docxtpl import DocxTemplate
from fastapi import HTTPException
from jinja2 import Environment, FileSystemLoader, select_autoescape
from weasyprint import HTML

from backend.app.core.paths import get_exports_dir
from backend.app.reporting.branding import BrandProfile, get_brand_profile

REPORTING_DIR = Path(__file__).resolve().parent.parent / "reporting"
TEMPLATE_PATH = REPORTING_DIR / "templates" / "Tender_Comparison_Template.docx"
PDF_TEMPLATE_NAME = "tender_comparison_report.html.j2"
ASSETS_DIR = REPORTING_DIR / "assets"
FONTS_DIR = ASSETS_DIR / "fonts"
TEMPLATES_DIR = REPORTING_DIR / "templates"


def _report_logo_path(brand: BrandProfile | None = None) -> Path | None:
    profile = brand or get_brand_profile()
    return profile.resolved_logo_path()


def _logo_src_for_pdf(logo_path: Path) -> str:
    """WeasyPrint base_url is REPORTING_DIR; prefer assets/ relative URLs when possible."""
    try:
        rel = logo_path.resolve().relative_to(REPORTING_DIR.resolve())
        return rel.as_posix()
    except ValueError:
        return logo_path.resolve().as_uri()


def _slug(value: str) -> str:
    cleaned = "".join(ch if ch.isalnum() or ch in ("-", "_") else "_" for ch in value)
    return cleaned.strip("_") or "report"


def _to_docx_text(value: Any) -> str:
    """Convert potential TinyMCE HTML into plain text safe for docxtpl placeholders."""
    text = str(value or "")
    if not text:
        return ""
    text = re.sub(r"</(p|div|li|h[1-6])\s*>", "\n", text, flags=re.IGNORECASE)
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", "", text)
    text = unescape(text)
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _clear_header_images(header) -> None:
    element = header._element
    for node in element.xpath('.//*[local-name()="drawing" or local-name()="pict"]'):
        parent = node.getparent()
        if parent is not None:
            parent.remove(node)


def _apply_docx_header_logo(docx_path: Path, logo_path: Path) -> None:
    if not logo_path.exists():
        return
    try:
        from docx import Document as DocxDocument
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Mm
    except ImportError:
        return
    doc = DocxDocument(str(docx_path))
    for section in doc.sections:
        header = section.header
        _clear_header_images(header)
        p = header.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run().add_picture(str(logo_path), width=Mm(42))
        body = header._element
        body.remove(p._element)
        body.insert(0, p._element)
    doc.save(str(docx_path))


def _prepare_docx_context(payload: dict[str, Any]) -> dict[str, Any]:
    report_context = (
        payload.get("report_context", {}) if isinstance(payload.get("report_context"), dict) else {}
    )
    draft_sections = (
        payload.get("draft_sections", {}) if isinstance(payload.get("draft_sections"), dict) else {}
    )
    project = (
        report_context.get("project", {}) if isinstance(report_context.get("project"), dict) else {}
    )
    commercial = (
        report_context.get("commercial", {})
        if isinstance(report_context.get("commercial"), dict)
        else {}
    )
    tender_meta = (
        report_context.get("tender_meta", {})
        if isinstance(report_context.get("tender_meta"), dict)
        else {}
    )
    executive = (
        draft_sections.get("executive_summary", {})
        if isinstance(draft_sections.get("executive_summary"), dict)
        else {}
    )
    commercial_section = (
        draft_sections.get("commercial_analysis", {})
        if isinstance(draft_sections.get("commercial_analysis"), dict)
        else {}
    )
    introduction_section = (
        draft_sections.get("introduction", {})
        if isinstance(draft_sections.get("introduction"), dict)
        else {}
    )

    next_steps = executive.get("next_steps", [])
    if not isinstance(next_steps, list):
        next_steps = []
    tender_rows = commercial.get("tender_comparison", [])
    if not isinstance(tender_rows, list):
        tender_rows = []
    first_tender_row = tender_rows[0] if tender_rows else {}
    tenderers = tender_meta.get("tenderers", [])
    if not isinstance(tenderers, list):
        tenderers = []

    contractor_names = [str(row.get("contractor") or "") for row in tender_rows]
    final_adjusted_values = [row.get("final_adjusted_tender_sum", 0) for row in tender_rows]
    construction_budget = commercial.get("construction_budget", 0)
    construction_budget_values = [construction_budget for _ in tender_rows]
    variance_values = [
        row.get("variance_to_construction_budget", row.get("variance_to_budget", 0))
        for row in tender_rows
    ]

    tender_review_rows = [
        {"label": "Final Adjusted Tender Sum", "values": final_adjusted_values},
        {"label": "Deduct Construction Budget", "values": construction_budget_values},
        {"label": "Variance to Construction Budget", "values": variance_values},
    ]
    return {
        "project_id": payload.get("project_id") or project.get("project_id") or "",
        "project_name": project.get("project_name") or "",
        "project_location": project.get("location") or "",
        "source_file_name": payload.get("source_file_name") or "",
        "project_description": project.get("project_description") or "",
        "responses_count": (report_context.get("tender_meta", {}) or {}).get("responses_count", ""),
        "executive_summary": _to_docx_text(executive.get("body") or ""),
        "recommendation": _to_docx_text(executive.get("recommendation") or ""),
        "next_steps": [_to_docx_text(step) for step in next_steps],
        "introduction": _to_docx_text(introduction_section.get("body") or ""),
        "tenderers": tenderers,
        "commercial_analysis": _to_docx_text(commercial_section.get("body") or ""),
        "construction_budget": construction_budget,
        "tender_rows": tender_rows,
        "row": first_tender_row,
        "rows": tender_rows,
        "tender_review_contractors": contractor_names,
        "tender_review_rows": tender_review_rows,
        "r": {"label": "", "values": []},
        "name": "",
        "v": "",
        "step": "",
    }


def _pdf_jinja_env() -> Environment:
    return Environment(
        loader=FileSystemLoader(str(TEMPLATES_DIR)),
        autoescape=select_autoescape(enabled_extensions=("html", "j2", "xml")),
    )


def _build_font_face_css(brand: BrandProfile) -> str:
    body = brand.font_body_file
    heading = brand.font_heading_file
    if not (FONTS_DIR / body).exists() or not (FONTS_DIR / heading).exists():
        return ""
    # Trusted brand-controlled font paths only — marked safe in template.
    family = brand.font_family.replace('"', "")
    body_safe = body.replace('"', "")
    heading_safe = heading.replace('"', "")
    return f"""
    @font-face {{
      font-family: "{family}";
      src: url("assets/fonts/{body_safe}") format("truetype");
      font-weight: 400;
      font-style: normal;
    }}
    @font-face {{
      font-family: "{family}";
      src: url("assets/fonts/{heading_safe}") format("truetype");
      font-weight: 700;
      font-style: normal;
    }}
"""


class ReportExporter(ABC):
    brand: BrandProfile

    def __init__(self, brand: BrandProfile | None = None) -> None:
        self.brand = brand or get_brand_profile()

    @abstractmethod
    def export(self, payload: dict[str, Any]) -> Path:
        raise NotImplementedError


class DocxExporter(ReportExporter):
    def export(self, payload: dict[str, Any]) -> Path:
        if not TEMPLATE_PATH.exists():
            raise HTTPException(
                status_code=500,
                detail=(
                    f"Word template not found at '{TEMPLATE_PATH}'. "
                    "Run: python -m backend.app.reporting.build_word_template"
                ),
            )
        context = _prepare_docx_context(payload)
        exports_dir = get_exports_dir()
        exports_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
        file_name = f"Tender_Comparison_{_slug(context['project_id'] or 'project')}_{stamp}.docx"
        output_path = exports_dir / file_name

        doc = DocxTemplate(str(TEMPLATE_PATH))
        doc.render(context)
        doc.save(str(output_path))
        logo_path = _report_logo_path(self.brand)
        if logo_path is not None:
            _apply_docx_header_logo(output_path, logo_path)
        return output_path


class PdfExporter(ReportExporter):
    def export(self, payload: dict[str, Any]) -> Path:
        exports_dir = get_exports_dir()
        exports_dir.mkdir(parents=True, exist_ok=True)
        project_id = payload.get("project_id") or "project"
        stamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
        file_name = f"Tender_Comparison_{_slug(str(project_id))}_{stamp}.pdf"
        output_path = exports_dir / file_name
        html = self._render_html(payload)
        HTML(string=html, base_url=str(REPORTING_DIR)).write_pdf(str(output_path))
        return output_path

    def _render_html(self, payload: dict[str, Any]) -> str:
        context = _prepare_docx_context(payload)
        brand = self.brand
        logo_path = brand.resolved_logo_path()
        logo_src = _logo_src_for_pdf(logo_path) if logo_path is not None else None
        template = _pdf_jinja_env().get_template(PDF_TEMPLATE_NAME)
        return template.render(
            **context,
            font_family=brand.font_family,
            accent_colour=brand.accent_colour,
            font_face_css=_build_font_face_css(brand),
            logo_src=logo_src,
            footer_lines=brand.pdf_footer_left_lines(),
            website=brand.website,
        )


def export_report_docx(payload: dict[str, Any], brand: BrandProfile | None = None) -> Path:
    return DocxExporter(brand=brand).export(payload)


def export_report_pdf(payload: dict[str, Any], brand: BrandProfile | None = None) -> Path:
    return PdfExporter(brand=brand).export(payload)
