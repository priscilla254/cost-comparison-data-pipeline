"""
Build the Word report template from BrandProfile + docxtpl placeholders.

Run from repo root:
  python -m backend.app.reporting.build_word_template

Re-run when branding or placeholder layout changes. Export fills this file;
it does not rebuild the template on every report.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from backend.app.reporting.branding import get_brand_profile

REPORTING_DIR = Path(__file__).resolve().parent
TEMPLATES_DIR = REPORTING_DIR / "templates"
OUTPUT_PATH = TEMPLATES_DIR / "Tender_Comparison_Template.docx"

# Secondary heading colour (matches PDF slate tone).
_SLATE = RGBColor(0x42, 0x56, 0x67)


def _hex_to_rgb(hex_colour: str) -> RGBColor:
    text = (hex_colour or "").strip().lstrip("#")
    if len(text) != 6:
        text = "32c3e2"
    return RGBColor(int(text[0:2], 16), int(text[2:4], 16), int(text[4:6], 16))


def _set_run_font(
    run, font_family: str, size_pt: float, *, bold: bool = False, colour: RGBColor | None = None
):
    run.font.name = font_family
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_family)
    run.font.size = Pt(size_pt)
    run.bold = bold
    if colour is not None:
        run.font.color.rgb = colour


def _add_heading(doc: Document, text: str, font_family: str, colour: RGBColor, size_pt: float = 14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    _set_run_font(run, font_family, size_pt, bold=True, colour=colour)
    return p


def _add_subheading(doc: Document, text: str, font_family: str, size_pt: float = 11):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    _set_run_font(run, font_family, size_pt, bold=True, colour=_SLATE)
    return p


def _add_body(doc: Document, text: str, font_family: str, size_pt: float = 10):
    """Add a paragraph whose text is a single run (safe for docxtpl tags)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    _set_run_font(run, font_family, size_pt)
    return p


def _add_label_value(doc: Document, label: str, placeholder: str, font_family: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    label_run = p.add_run(f"{label}: ")
    _set_run_font(label_run, font_family, 10, bold=True)
    value_run = p.add_run(placeholder)
    _set_run_font(value_run, font_family, 10)
    return p


def _configure_header_footer(doc: Document, brand, accent: RGBColor, font_family: str) -> None:
    section = doc.sections[0]
    section.top_margin = Mm(28)
    section.bottom_margin = Mm(28)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    logo = brand.resolved_logo_path()
    if logo is not None:
        hp.add_run().add_picture(str(logo), width=Mm(42))
    else:
        run = hp.add_run(brand.company_name)
        _set_run_font(run, font_family, 9, bold=True, colour=accent)

    footer = section.footer
    footer.is_linked_to_previous = False
    # Remove default empty paragraph(s), then add rule + brand lines.
    for p in list(footer.paragraphs):
        p._element.getparent().remove(p._element)

    rule_p = footer.add_paragraph()
    pPr = rule_p._p.get_or_add_pPr()
    pBdr = pPr.makeelement(qn("w:pBdr"), {})
    bottom = pBdr.makeelement(
        qn("w:bottom"),
        {
            qn("w:val"): "single",
            qn("w:sz"): "12",
            qn("w:space"): "4",
            qn("w:color"): str(brand.accent_colour).lstrip("#").upper(),
        },
    )
    pBdr.append(bottom)
    pPr.append(pBdr)

    left = footer.add_paragraph()
    left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    lines = brand.pdf_footer_left_lines()
    if lines:
        run = left.add_run("\n".join(lines))
        _set_run_font(run, font_family, 8, colour=_SLATE)

    center = footer.add_paragraph()
    center.alignment = WD_ALIGN_PARAGRAPH.CENTER
    web_run = center.add_run(brand.website)
    _set_run_font(web_run, font_family, 9, bold=True, colour=_SLATE)


def build_template(output_path: Path | None = None) -> Path:
    brand = get_brand_profile()
    font_family = brand.font_family or "Arial"
    accent = _hex_to_rgb(brand.accent_colour)
    dest = output_path or OUTPUT_PATH
    dest.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = font_family
    style.font.size = Pt(10)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), font_family)

    _configure_header_footer(doc, brand, accent, font_family)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(8)
    title_run = title.add_run("Tender Comparison Report")
    _set_run_font(title_run, font_family, 18, bold=True, colour=_SLATE)

    _add_label_value(doc, "Project ID", "{{ project_id }}", font_family)
    _add_label_value(doc, "Project Name", "{{ project_name }}", font_family)
    _add_label_value(doc, "Location", "{{ project_location }}", font_family)

    _add_heading(doc, "01 - Executive Summary", font_family, accent, size_pt=12)
    _add_body(doc, "{{ executive_summary }}", font_family)

    _add_subheading(doc, "Project Information", font_family)
    _add_label_value(doc, "Project Description", "{{ project_description }}", font_family)
    _add_label_value(doc, "Number of Responses", "{{ responses_count }}", font_family)
    _add_label_value(doc, "Tenders Issued", "{{ tenders_issued_date }}", font_family)
    _add_label_value(doc, "Tender Deadline", "{{ tender_deadline_date }}", font_family)
    _add_label_value(doc, "Addendums Issued", "{{ addendums_issued_count }}", font_family)

    _add_subheading(doc, "Tender Review", font_family)
    _add_body(
        doc,
        "Item | {% for name in tender_review_contractors %}{{ name }} | {% endfor %}",
        font_family,
    )
    _add_body(
        doc,
        "{% for r in tender_review_rows %}{{ r.label }} | {% for v in r.values %}{{ v }} | {% endfor %}\n{% endfor %}",
        font_family,
    )

    _add_subheading(doc, "Recommendation", font_family)
    _add_body(doc, "{{ recommendation }}", font_family)

    _add_subheading(doc, "Recommended Next Steps", font_family)
    _add_body(doc, "{% for step in next_steps %}- {{ step }}\n{% endfor %}", font_family)

    _add_heading(doc, "02 - Introduction", font_family, accent, size_pt=12)
    _add_subheading(doc, "Report Overview", font_family)
    _add_body(doc, "{{ introduction }}", font_family)

    _add_subheading(doc, "Tenderer List", font_family)
    _add_body(doc, "{% for name in tenderers %}- {{ name }}\n{% endfor %}", font_family)

    _add_heading(doc, "03 - Commercial Analysis", font_family, accent, size_pt=12)
    _add_body(doc, "{{ commercial_analysis }}", font_family)

    _add_subheading(doc, "Tender Comparison", font_family)
    _add_body(
        doc,
        "{% for row in tender_rows %}{{ row.contractor }} | {{ row.final_adjusted_tender_sum }}\n{% endfor %}",
        font_family,
    )

    doc.save(str(dest))
    return dest


def main() -> None:
    path = build_template()
    brand = get_brand_profile()
    print(f"Wrote Word template: {path}")
    print(f"  company={brand.company_name!r} accent={brand.accent_colour!r}")
    print(f"  logo={'yes' if brand.resolved_logo_path() else 'no'}")


if __name__ == "__main__":
    main()
