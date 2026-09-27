"""
Build a minimal workbook that passes sheet discovery but fails Python validation.

Missing required ProjectInformation columns after normalize → MISSING_COLUMN /
ROW_COUNT errors → BatchStatus FAILED.

Output: sample_data/DEMO_Failing_Workbook.xlsx
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "sample_data" / "DEMO_Failing_Workbook.xlsx"


def build_failing_workbook(path: Path = OUTPUT) -> Path:
    wb = Workbook()
    # Default sheet renamed into Project Information alias with incomplete fields.
    ws_pi = wb.active
    ws_pi.title = "Project Information - 1"
    ws_pi.append(["Field", "Value"])
    ws_pi.append(["Project Name", "Broken Demo"])
    # Intentionally omit Project Number / Location / Sector / Selected contractor.

    ws_pq = wb.create_sheet("Project Information - 2")
    ws_pq.append(["Code", "Name", "Qty", "Unit"])
    ws_pq.append(["PQ-1", "GIFA", "not-a-number", "m2"])

    ws_eq = wb.create_sheet("Project Information - 3")
    ws_eq.append(["Code", "Element", "Qty", "Unit"])
    ws_eq.append(["1.1", "Foundations", 10, "m2"])

    ws_sum = wb.create_sheet("SUMMARY")
    ws_sum.append(["Ref", "Element", "Rate", "Total"])
    ws_sum.append([1.0, "Substructure", None, None])
    ws_sum.append([1.1, "Foundations", 10, 100])

    ws_l3 = wb.create_sheet("1.1 Foundations")
    ws_l3.append([None, None, "Qty", "Unit", "Rate", "Total"])
    ws_l3.append(["x", "Dig", 1, "m3", 10, 10])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def main() -> None:
    out = build_failing_workbook()
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
